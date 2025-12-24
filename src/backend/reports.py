"""
Módulo de Relatórios - PontoPrime V2
Geração de relatórios em PDF e Excel
"""
from datetime import datetime
from typing import List
from io import BytesIO
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, Image
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from models import PunchRecord, Employee


# ========================================
# PDF REPORTS
# ========================================

def generate_punch_records_pdf(
    records: List[tuple],  # (PunchRecord, employee_name, employee_code)
    title: str = "Relatório de Registros de Ponto"
) -> BytesIO:
    """
    Gera PDF com registros de ponto

    Args:
        records: Lista de tuplas (PunchRecord, employee_name, employee_code)
        title: Título do relatório

    Returns:
        BytesIO contendo o PDF
    """
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(A4))
    elements = []

    # Styles
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'CustomTitle',
        parent=styles['Heading1'],
        fontSize=18,
        textColor=colors.HexColor('#0066A1'),
        alignment=TA_CENTER,
        spaceAfter=30
    )

    # Título
    elements.append(Paragraph(title, title_style))
    elements.append(Spacer(1, 0.5*cm))

    # Info
    info_text = f"Gerado em: {datetime.now().strftime('%d/%m/%Y às %H:%M:%S')}"
    elements.append(Paragraph(info_text, styles['Normal']))
    elements.append(Spacer(1, 0.5*cm))

    # Tabela de dados
    data = [['ID', 'Código', 'Funcionário', 'Data/Hora', 'Localização', 'Biometria', 'Status']]

    for record, employee_name, employee_code in records:
        location = ""
        if record.latitude and record.longitude:
            location = f"{record.latitude:.4f}, {record.longitude:.4f}"

        bio = "✓" if record.biometric_verified else "✗"

        data.append([
            str(record.id),
            employee_code,
            employee_name,
            record.timestamp.strftime('%d/%m/%Y %H:%M'),
            location,
            bio,
            record.status
        ])

    # Criar tabela
    table = Table(data, colWidths=[2*cm, 3*cm, 6*cm, 4*cm, 5*cm, 2*cm, 3*cm])

    # Estilo da tabela
    table.setStyle(TableStyle([
        # Header
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0066A1')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 10),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 12),

        # Body
        ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
        ('TEXTCOLOR', (0, 1), (-1, -1), colors.black),
        ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
        ('FONTSIZE', (0, 1), (-1, -1), 8),
        ('GRID', (0, 0), (-1, -1), 1, colors.black),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),

        # Alternating rows
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.lightgrey]),
    ]))

    elements.append(table)

    # Footer
    elements.append(Spacer(1, 1*cm))
    footer_text = f"Total de registros: {len(records)}"
    elements.append(Paragraph(footer_text, styles['Normal']))

    # Build PDF
    doc.build(elements)
    buffer.seek(0)

    return buffer


# ========================================
# EXCEL REPORTS
# ========================================

def generate_punch_records_excel(
    records: List[tuple],  # (PunchRecord, employee_name, employee_code)
    title: str = "Relatório de Registros de Ponto"
) -> BytesIO:
    """
    Gera Excel com registros de ponto

    Args:
        records: Lista de tuplas (PunchRecord, employee_name, employee_code)
        title: Título do relatório

    Returns:
        BytesIO contendo o Excel
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Registros de Ponto"

    # Styles
    header_fill = PatternFill(start_color="0066A1", end_color="0066A1", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=12)
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    # Título
    ws.merge_cells('A1:H1')
    title_cell = ws['A1']
    title_cell.value = title
    title_cell.font = Font(bold=True, size=16, color="0066A1")
    title_cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 30

    # Info
    ws.merge_cells('A2:H2')
    info_cell = ws['A2']
    info_cell.value = f"Gerado em: {datetime.now().strftime('%d/%m/%Y às %H:%M:%S')}"
    info_cell.alignment = Alignment(horizontal='left')
    ws.row_dimensions[2].height = 20

    # Headers
    headers = ['ID', 'Código Func.', 'Funcionário', 'Data', 'Hora', 'Latitude', 'Longitude', 'Biometria', 'Status']
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col_num)
        cell.value = header
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = border

    # Data
    for row_num, (record, employee_name, employee_code) in enumerate(records, 5):
        ws.cell(row=row_num, column=1).value = record.id
        ws.cell(row=row_num, column=2).value = employee_code
        ws.cell(row=row_num, column=3).value = employee_name
        ws.cell(row=row_num, column=4).value = record.timestamp.strftime('%d/%m/%Y')
        ws.cell(row=row_num, column=5).value = record.timestamp.strftime('%H:%M:%S')
        ws.cell(row=row_num, column=6).value = float(record.latitude) if record.latitude else ""
        ws.cell(row=row_num, column=7).value = float(record.longitude) if record.longitude else ""
        ws.cell(row=row_num, column=8).value = "Sim" if record.biometric_verified else "Não"
        ws.cell(row=row_num, column=9).value = record.status

        # Apply border
        for col_num in range(1, 10):
            ws.cell(row=row_num, column=col_num).border = border

    # Adjust column widths
    column_widths = [8, 12, 25, 12, 10, 12, 12, 12, 12]
    for col_num, width in enumerate(column_widths, 1):
        ws.column_dimensions[get_column_letter(col_num)].width = width

    # Footer
    footer_row = len(records) + 6
    ws.merge_cells(f'A{footer_row}:I{footer_row}')
    footer_cell = ws.cell(row=footer_row, column=1)
    footer_cell.value = f"Total de registros: {len(records)}"
    footer_cell.font = Font(bold=True)

    # Save to buffer
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    return buffer
